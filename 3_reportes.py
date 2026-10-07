# -*- coding: utf-8 -*-
"""
PASO 5 · REPORTES  (módulos 9, 10, 12 y 13: riesgo residual, seguimiento MEA, explicación y trazabilidad)

Lee  salidas/preparado/  y  salidas/motor/  y escribe  salidas/reportes/:
  portafolio_final.csv        Producto 2 · qué, dónde, cuánto, actores, orden, beneficio, cambio con SSP3-7.0/2060
  riesgo_residual.csv         Producto 3 · qué queda sin resolver (municipio × dimensión)
  seguimiento_MEA.csv         Producto 3 · indicadores de actividad, resultado y vulnerabilidad
  brechas_informacion.csv     Producto 3 · qué datos y dependencias debe levantar CORNARE
  resultados.json             todo junto, para el tablero (frontend Vue)
  reporte_completo.xlsx       todo en una hoja de cálculo para el jurado

Uso:   python 3_reportes.py
"""
import json
import os
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

AQUI = os.path.dirname(os.path.abspath(__file__))
PREP = os.path.join(AQUI, "salidas", "preparado")
MOT = os.path.join(AQUI, "salidas", "motor")
SAL = os.path.join(AQUI, "salidas", "reportes")
MUNICIPIOS = ["Rionegro", "Guarne", "Marinilla"]
DIMENSIONES = ["Biodiversidad", "Recurso hídrico", "Seguridad alimentaria", "Hábitat",
               "Infraestructura", "Riesgo de desastres", "Salud"]

# ----------------------------------------------------------------------------------------------
# Textos por medida. Cada texto lleva su etiqueta de trazabilidad.
#   reduce / fortalece : [INFERENCIA] a partir de la unidad de intervención del reto
#   actores            : [OFICIAL] si la fuente lo nombra; [INFERENCIA] si sale de competencias
#   resultado          : [PROPUESTO] indicador de resultado sugerido por el equipo
# ----------------------------------------------------------------------------------------------
TEXTOS = {
    "M01": dict(reduce="Pérdida de cobertura en áreas estratégicas", fortalece="Instrumento económico de conservación con propietarios",
                actores="CORNARE y alcaldía (PSA registrado en el PICCA) [OFICIAL: Excel adaptación]; propietarios rurales [INFERENCIA]",
                resultado="Hectáreas conservadas bajo PSA en áreas estratégicas"),
    "M02": dict(reduce="Suelos desprotegidos y ecosistemas degradados (cobeneficio en laderas)", fortalece="—",
                actores="CORNARE (restauración registrada en el PICCA) [OFICIAL: Excel]; alcaldía; comunidades y JAC [INFERENCIA]",
                resultado="Hectáreas restauradas con sobrevivencia a 3 años; nº de movimientos en masa en las zonas intervenidas"),
    "M03": dict(reduce="Fragmentación de ecosistemas (conectividad)", fortalece="Gestión, acuerdos y figuras de conservación en el corredor",
                actores="CORNARE como autoridad ambiental [OFICIAL: Reto]; alcaldías (ordenamiento) y propietarios [INFERENCIA]",
                resultado="Hectáreas con figura de conservación o acuerdo vigente"),
    "M04": dict(reduce="Presión sobre el recurso hídrico por demanda", fortalece="Medición y gestión del agua en grandes usuarios",
                actores="CORNARE; empresas ACV [OFICIAL: Reto sec. 2]; empresas de servicios públicos [INFERENCIA]",
                resultado="m³ de agua ahorrados o reusados por los usuarios intervenidos"),
    "M05": dict(reduce="Ocupación de retiros y llanuras de inundación", fortalece="—",
                actores="CORNARE y Secretaría de Planeación municipal (control de ocupación de retiros) [OFICIAL: recomendaciones 2014]; propietarios de predios en llanuras [OFICIAL: 2014]",
                resultado="Hectáreas de ronda recuperadas; predios o viviendas en retiro; Cambio en población urbana localizada en zonas de amenaza [OFICIAL: indicador Excel]"),
    "M06": dict(reduce="Degradación de ecosistemas abastecedores", fortalece="—",
                actores="CORNARE (POMCA) [OFICIAL: Zonificación Guarne 2014]; acueductos [INFERENCIA]",
                resultado="Hectáreas protegidas en la microcuenca abastecedora"),
    "M07": dict(reduce="Prácticas agrícolas que degradan el suelo (cultivos limpios, sobrepastoreo)", fortalece="Asistencia técnica a productores",
                actores="Alcaldía (agricultura) [INFERENCIA]; productores; JAC y acueductos rurales [OFICIAL: recomendación Marinilla 2014]; CORNARE",
                resultado="Hectáreas con prácticas de conservación de suelo; nº de movimientos en masa en las veredas intervenidas"),
    "M08": dict(reduce="Manejo de suelo y agua en fincas", fortalece="Asistencia técnica y asociatividad",
                actores="Alcaldía y familias campesinas (producción agroecológica registrada) [OFICIAL: Excel]; CORNARE",
                resultado="Familias con prácticas productivas adaptadas"),
    "M09": dict(reduce="Calor urbano y escorrentía", fortalece="—", actores="Alcaldía; CORNARE [INFERENCIA]",
                resultado="Hectáreas de área verde multifuncional"),
    "M10": dict(reduce="Escorrentía urbana e inundación por drenaje", fortalece="—",
                actores="Alcaldía (planeación) y empresa de alcantarillado [INFERENCIA]; CORNARE",
                resultado="m³ de escorrentía retenida; puntos de inundación urbana; Cambio en población urbana localizada en zonas de amenaza [OFICIAL: indicador Excel]"),
    "M11": dict(reduce="Exposición de un punto crítico de vía o servicio", fortalece="—",
                actores="Alcaldía y entidad responsable de la vía o servicio [INFERENCIA]; consejo municipal de gestión del riesgo",
                resultado="Población y servicios protegidos en el punto crítico; Cambio en población urbana localizada en zonas de amenaza [OFICIAL: indicador Excel]"),
    "M12": dict(reduce="Interrupción de servicios esenciales", fortalece="—", actores="Empresas de servicios públicos; alcaldía [INFERENCIA]",
                resultado="Nº de interrupciones del servicio por eventos climáticos"),
    "M13": dict(reduce="—", fortalece="Anticipación y respuesta: instrumentación, comunicaciones y protocolos",
                actores="CORNARE; consejos municipales de gestión del riesgo [OFICIAL: Marinilla 2014 / INFERENCIA]; comunidades (diálogo recomendado en Guarne 2012) [OFICIAL]",
                resultado="Puntos críticos con alerta operativa; población cubierta; tiempo de aviso"),
    "M14": dict(reduce="—", fortalece="Conocimiento del riesgo: estudios, modelación, protocolos y plataforma",
                actores="CORNARE (Observatorio, gestión del riesgo) [OFICIAL]; consejos municipales de gestión del riesgo [INFERENCIA]; empresas ACV con información agregada [OFICIAL: Reto]",
                resultado="% de celdas de la ficha V/CA con dato oficial (línea base calculada abajo); nº de dependencias empresa–territorio caracterizadas"),
    "M15": dict(reduce="—", fortalece="Vigilancia climato-sensible y capacidad de respuesta en salud", actores="Secretarías de salud [INFERENCIA]; CORNARE",
                resultado="Casos de enfermedades climato-sensibles notificados y atendidos"),
}
ORDEN_TIPO = {"M14": 0}                          # [SUPUESTO] primero lo que produce información para decidir lo demás
RANGO_TIPO = {"CA": 1, "S+CA": 2, "S": 3}        # [SUPUESTO] luego capacidad (gobernanza), después obras sobre sensibilidad


def cargar():
    for f in ("resumen.json", "alternativas.csv"):
        if not os.path.exists(os.path.join(MOT, f)):
            sys.exit("Primero ejecute: python 1_preparar.py  y  python 2_motor.py")
    rp = lambda f: pd.read_csv(os.path.join(PREP, f), encoding="utf-8-sig")
    rm = lambda f: pd.read_csv(os.path.join(MOT, f), encoding="utf-8-sig")
    d = dict(ficha=rp("ficha_larga.csv"), medidas=rp("medidas.csv"), evidencia=rp("evidencia.csv"), pesos=rp("pesos.csv"),
             alts=rm("alternativas.csv"), regret=rm("regret_portafolios.csv"), cf=rm("contrafactuales.csv"),
             vacias=rm("celdas_vacias_usadas.csv"),
             resumen=json.load(open(os.path.join(MOT, "resumen.json"), encoding="utf-8")),
             par=json.load(open(os.path.join(PREP, "parametros.json"), encoding="utf-8")))
    vf = os.path.join(MOT, "datos_que_cambian_la_decision.csv")
    d["voi"] = pd.read_csv(vf, encoding="utf-8-sig") if os.path.getsize(vf) > 5 else pd.DataFrame()
    hf = os.path.join(PREP, "historico_corredor.csv")
    d["hist"] = pd.read_csv(hf, encoding="utf-8-sig") if os.path.exists(hf) else pd.DataFrame()
    return d


def fmt(v):
    return "NO DISPONIBLE" if v is None or pd.isna(v) else f"{v:.2f}".replace(".", ",")


def valor(ficha, m, d, comp):
    r = ficha[(ficha.municipio == m) & (ficha.dimension == d) & (ficha.componente == comp)]
    if r.empty or r.origen.iloc[0] == "NO DISPONIBLE":
        return None, "NO DISPONIBLE"
    return float(r.valor.iloc[0]), r.origen.iloc[0]


# ----------------------------------------------------------------------------------------------
# DECISIÓN FINAL según el parámetro A / B / C
# ----------------------------------------------------------------------------------------------
def portafolio_final(d):
    res, dec = d["resumen"], d["par"].get("decision_stress", "A")
    if dec == "B":
        return res["stress_2060"], "B · se adopta el portafolio del stress test SSP3-7.0/2060"
    if dec == "C":
        reg = d["regret"]; eq = reg[reg.regret_p90 <= reg.regret_p90.min() + 0.01]
        nuevas = set(res["entra_2060"])
        con = eq[eq.portafolio.apply(lambda p: bool(nuevas & set(p.split(" + "))))].sort_values("regret_p90")
        if len(con):
            return con.portafolio.iloc[0].split(" + "), "C · híbrido: portafolio casi equivalente que incorpora una medida del stress test"
        return res["robusto"], "C solicitado, pero ningún portafolio equivalente incluye medidas del stress test → se usa el robusto"
    return res["robusto"], "A · se mantiene el portafolio robusto (mínimo arrepentimiento frente a hoy y 2060)"


# ----------------------------------------------------------------------------------------------
# MÓDULO 12 · EXPLICACIÓN + PRODUCTO 2
# ----------------------------------------------------------------------------------------------
def construir_portafolio(d, elegidas):
    alts, med, evid, ficha, cf, res = d["alts"], d["medidas"], d["evidencia"], d["ficha"], d["cf"], d["resumen"]
    compiten = alts[alts.elegible].sort_values("puntaje_ref", ascending=False).reset_index(drop=True)
    vca = ficha[ficha.componente.isin(["Vulnerabilidad", "Capacidad adaptativa"])]
    base_pct = (vca.origen != "NO DISPONIBLE").mean()
    filas = []
    for alt in elegidas:
        a = alts[alts.alternativa == alt].iloc[0]; mid = alt.split("-")[0]; m = med[med.ID == mid].iloc[0]
        munis = MUNICIPIOS if a.territorio == "Corredor" else [a.territorio]
        dim = a.dimension
        # problema: datos oficiales de la dimensión + problema documentado
        datos = []
        for x in munis:
            v, ov = valor(ficha, x, dim, "Vulnerabilidad"); c, oc = valor(ficha, x, dim, "Capacidad adaptativa")
            datos.append(f"{x}: V {fmt(v)}{' [' + ov + ']' if v is not None else ''}, CA {fmt(c)}{' [' + oc + ']' if c is not None else ''}")
        ev = evid[(evid.ID == mid) & (evid.Municipio.isin(munis)) & (evid.Evidencia > 0)]
        localiz = " | ".join(f"{r.Municipio}: {r['Qué dice la fuente']} ({r.Fuente})" for _, r in ev.iterrows()) or "Sin localización documentada en las fuentes"
        # explicación: posición, criterios, estabilidad, qué la reemplaza
        pos = int(compiten.index[compiten.alternativa == alt][0]) + 1 if alt in set(compiten.alternativa) else None
        f_med = res["frecuencia_por_medida"].get(mid, 0)
        sin = cf[cf.caso == f"Sin {alt}"]
        reemplazo = ""
        if len(sin):
            nuevo = set(sin.portafolio.iloc[0].split(" + ")) - set(elegidas)
            reemplazo = f"Si se elimina, el modelo la reemplaza por: {', '.join(sorted(nuevo)) or '—'}."
        por_que = [f"Puntaje {a.puntaje_ref:.3f} (puesto {pos} de {len(compiten)} alternativas que compiten)",
                   f"Necesidad {a.Necesidad:.2f} · Brecha CA {a['Brecha CA']:.2f} · Evidencia {a.Evidencia:.2f} · "
                   f"Cobertura {a.Cobertura:.2f} · Cobeneficios {a.Cobeneficios:.2f}",
                   f"La medida aparece en el {f_med:.0%} de los 3.000 escenarios simulados; confianza {a.confianza}",
                   reemplazo]
        # cambio con el stress test
        if alt in res["sale_2060"]:
            cambio = "Sale en el portafolio SSP3-7.0/2060 (se mantiene por robustez frente a hoy y 2060)"
        elif alt in res["entra_2060"]:
            cambio = "Entra con el stress test SSP3-7.0/2060"
        else:
            cambio = "Se mantiene en SSP3-7.0/2060"
        cambio += f" · Robustez: {a.robustez_2060}"
        tx = TEXTOS[mid]
        tipo = ORDEN_TIPO.get(mid, RANGO_TIPO.get(a.actua_sobre, 3))
        filas.append(dict(
            alternativa=alt, id=mid, medida=m.Medida, dimension=dim, territorio=a.territorio, costo_M=float(a.costo_M),
            problema=f"{dim} — " + "; ".join(datos),
            sensibilidad_que_reduce=tx["reduce"], capacidad_que_fortalece=tx["fortalece"], localizacion=localiz,
            actores=tx["actores"], beneficio_esperado=(
                f"Reduce: {tx['reduce']}. Fortalece: {tx['fortalece']}. Magnitud: no cuantificable con los datos disponibles "
                f"(efectividad NO DISPONIBLE)."),
            explicacion=" · ".join(p for p in por_que if p), cambio_ssp370_2060=cambio,
            indicador_actividad=f"{m['Indicador MEA disponible']} [{m['Estado del indicador']}]",
            indicador_resultado=f"{tx['resultado'].replace('(línea base calculada abajo)', f'(línea base {base_pct:.0%})')} [PROPUESTO]",
            confianza=a.confianza,
            _tipo=tipo, _puntaje=float(a.puntaje_ref)))
    df = pd.DataFrame(filas).sort_values(["_tipo", "_puntaje"], ascending=[True, False]).reset_index(drop=True)
    df.insert(0, "orden", range(1, len(df) + 1))
    df["regla_de_orden"] = "[SUPUESTO] 1º información (M14) → 2º capacidad adaptativa → 3º mixtas → 4º sensibilidad; empate: mayor puntaje"
    return df.drop(columns=["_tipo", "_puntaje"])


# ----------------------------------------------------------------------------------------------
# MÓDULO 9 · RIESGO RESIDUAL (municipio × dimensión)
# ----------------------------------------------------------------------------------------------
def construir_residual(d, port):
    ficha, med, evid = d["ficha"], d["medidas"], d["evidencia"]
    cob = {r.ID: [c.strip() for c in str(r["Cobeneficios (otras dimensiones)"]).split(";") if c.strip() and c.strip() != "nan"]
           for _, r in med.iterrows()}
    cubiertas = set()                    # (medida, municipio) cubiertos por una alternativa elegida
    for _, p in port.iterrows():
        for x in (MUNICIPIOS if p.territorio == "Corredor" else [p.territorio]):
            cubiertas.add((p.id, x))
    filas = []
    for x in MUNICIPIOS:
        for dim in DIMENSIONES:
            directas, indirectas, corredor = [], [], []
            for _, p in port.iterrows():
                aplica = p.territorio in (x, "Corredor")
                if not aplica:
                    continue
                if p.dimension == dim:
                    (corredor if p.territorio == "Corredor" else directas).append(p.alternativa)
                elif dim in cob.get(p.id, []):
                    indirectas.append(p.alternativa)
            sit = []
            for comp in ("Amenaza", "Sensibilidad", "Capacidad adaptativa", "Vulnerabilidad", "Riesgo hoy", "Riesgo SSP3-7.0 2060"):
                v, o = valor(ficha, x, dim, comp)
                if v is not None:
                    sit.append(f"{comp} {fmt(v)} [{o}]")
            # problemas documentados (evidencia ≥ 0,5) de medidas de esta dimensión que NO quedaron en el portafolio
            ids_dim = med[med["Dimensión"] == dim].ID.tolist()
            pend = evid[(evid.Municipio == x) & (evid.ID.isin(ids_dim)) & (evid.Evidencia >= 0.5)]
            pend = pend[pend.ID.map(lambda i: (i, x) not in cubiertas).astype(bool)] if len(pend) else pend
            fuerte = (pend.Evidencia >= 1).any()
            pend_txt = " | ".join(f"{r.ID} (evidencia {fmt(r.Evidencia)}): {r['Qué dice la fuente']}" for _, r in pend.iterrows())
            # [SUPUESTO] clasificación de cobertura de la intervención
            if directas:
                estado = "Atendido con medida municipal directa"
            elif corredor:
                estado = "Parcial: solo medida de corredor"
            elif indirectas:
                estado = "Parcial: solo cobeneficio de otra medida"
            else:
                estado = "Sin atender"
            v, _ = valor(ficha, x, dim, "Vulnerabilidad")
            if v is None and not sit:
                gravedad = "DESCONOCIDA (sin datos: no se puede afirmar que sea baja)"
            elif v is not None and v >= 0.6:
                gravedad = "Vulnerabilidad alta documentada"
            else:
                gravedad = "Sin vulnerabilidad alta documentada"
            filas.append(dict(municipio=x, dimension=dim, situacion_inicial="; ".join(sit) or "NO DISPONIBLE",
                              medidas_directas=", ".join(directas), medidas_corredor=", ".join(corredor),
                              medidas_por_cobeneficio=", ".join(indirectas), estado_intervencion=estado,
                              gravedad_conocida=gravedad,
                              reduccion_esperada="No cuantificable con los datos disponibles",
                              problemas_documentados_sin_medida=pend_txt or "—",
                              # [SUPUESTO] ALTA: sin medida municipal directa y (V ≥ 0,6 documentada o problema con evidencia 1);
                              # MEDIA: problema con evidencia 0,5 sin medida, o solo cobertura parcial; DESCONOCIDA: sin datos ni medida
                              prioridad_residual=("ALTA" if (not directas and (gravedad.startswith("Vulnerabilidad alta") or fuerte)) else
                                                  "MEDIA" if (not directas and (pend_txt or estado.startswith("Parcial"))) else
                                                  "DESCONOCIDA" if gravedad.startswith("DESCONOCIDA") and estado == "Sin atender" else
                                                  "BAJA" if directas else "SIN PROBLEMA DOCUMENTADO")))
    return pd.DataFrame(filas)


# ----------------------------------------------------------------------------------------------
# MÓDULO 10 · SEGUIMIENTO MEA
# ----------------------------------------------------------------------------------------------
def construir_mea(d, port):
    ficha = d["ficha"]
    vca = ficha[ficha.componente.isin(["Vulnerabilidad", "Capacidad adaptativa"])]
    base_pct = (vca.origen != "NO DISPONIBLE").mean()
    filas = []
    for _, p in port.iterrows():
        munis = MUNICIPIOS if p.territorio == "Corredor" else [p.territorio]
        lineas = []
        for x in munis:
            v, _ = valor(ficha, x, p.dimension, "Vulnerabilidad"); c, _ = valor(ficha, x, p.dimension, "Capacidad adaptativa")
            lineas.append(f"{x}: V {fmt(v)} · CA {fmt(c)}")
        res_txt = p.indicador_resultado.replace("(línea base calculada abajo)", f"(línea base {base_pct:.0%})")
        filas.append(dict(orden=p.orden, alternativa=p.alternativa, medida=p.medida,
                          actividad_que_se_hizo=p.indicador_actividad,
                          resultado_que_cambio=res_txt,
                          vulnerabilidad_se_redujo=f"Índices CORNARE de {p.dimension} — línea base: " + "; ".join(lineas),
                          periodicidad=("Actividad y resultado: anual [los indicadores del Excel se reportan anualmente]; "
                                        "índices V/CA: en cada actualización del estudio CORNARE"),
                          responsable_reporte=p.actores.split(";")[0]))
    return pd.DataFrame(filas), base_pct


# ----------------------------------------------------------------------------------------------
# BRECHAS DE INFORMACIÓN
# ----------------------------------------------------------------------------------------------
def construir_brechas(d, base_pct):
    ficha, voi = d["ficha"], d["voi"]
    filas = []
    for comp in ("Vulnerabilidad", "Capacidad adaptativa", "Riesgo SSP3-7.0 2060", "Sensibilidad", "Amenaza"):
        sub = ficha[ficha.componente == comp]; n = (sub.origen != "NO DISPONIBLE").sum()
        filas.append(dict(tipo="Ficha CORNARE", dato=f"{comp} por municipio × dimensión", disponible=f"{n} de {len(sub)}",
                          por_que_importa="Base de los criterios del modelo" if comp in ("Vulnerabilidad", "Capacidad adaptativa")
                          else "Stress test SSP3-7.0/2060" if "2060" in comp else "Contexto (no entra al puntaje)",
                          umbral_que_cambia_la_decision="Ver filas 'Dato que cambia la decisión'"))
    for _, r in voi.iterrows() if len(voi) else []:
        partes = []
        for col, txt in (("cambia_QUE_si_sube_a", "cambia QUÉ medidas si ≥"), ("cambia_QUE_si_baja_a", "cambia QUÉ medidas si ≤"),
                         ("cambia_DONDE_si_sube_a", "cambia DÓNDE si ≥"), ("cambia_DONDE_si_baja_a", "cambia DÓNDE si ≤")):
            if col in r and not pd.isna(r[col]):
                partes.append(f"{txt} {r[col]:.2f}")
        filas.append(dict(tipo="Dato que cambia la decisión", dato=r.celda_vacia, disponible="NO",
                          por_que_importa=f"Resultado si se cruza el umbral: {r.portafolio_resultante}",
                          umbral_que_cambia_la_decision="; ".join(partes)))
    fijas = [
        ("Dependencias (Reto sec. 5)", "Empresa ACV → captación o fuente hídrica, caudal concesionado vs usado", "NO",
         "Sin esto no se puede priorizar 'Uso intersectorial eficiente del agua' (M04)"),
        ("Dependencias (Reto sec. 5)", "Empresa ACV → vía de acceso crítica, circuito eléctrico, proveedores críticos de la región, trabajadores por municipio", "NO",
         "Sin esto no se puede priorizar infraestructura (M11, M12) ni saber qué interrupción se propaga"),
        ("Dependencias (Reto sec. 5)", "Punto crítico de vía o servicio → población y empresas que dependen, rutas alternas, tiempo tolerable de interrupción", "NO",
         "Define dónde una obra de infraestructura protege más"),
        ("Efectividad", "Línea base y medición posterior del indicador de resultado de cada medida", "NO",
         "Permitiría estimar efectividad (hoy α se prueba de 0 a 1) y, a futuro, entrenar un modelo de aprendizaje automático"),
        ("Salud", "Vigilancia de enfermedades climato-sensibles por municipio", "NO", "La medida de salud (M15) no compite por falta de datos"),
        ("Registro histórico", "Medidas de adaptación de Marinilla (1 registro 2018–2025)", "Incompleto", "Subregistro: no indica menor necesidad"),
    ]
    for t, dat, disp, por in fijas:
        filas.append(dict(tipo=t, dato=dat, disponible=disp, por_que_importa=por, umbral_que_cambia_la_decision="—"))
    return pd.DataFrame(filas)


# ----------------------------------------------------------------------------------------------
# HALLAZGOS (máximo 5, calculados a partir de los datos)
# ----------------------------------------------------------------------------------------------
def construir_hallazgos(d, port, residual, base_pct):
    ficha, res = d["ficha"], d["resumen"]
    H = []
    v = ficha[(ficha.componente == "Vulnerabilidad") & (ficha.origen != "NO DISPONIBLE")].sort_values("valor", ascending=False)
    if len(v):
        top = v.head(2)
        txt = " y ".join(f"{r.dimension} en {r.municipio} ({fmt(r.valor)})" for r in top.itertuples())
        H.append(dict(hallazgo=f"La vulnerabilidad más alta con dato es {txt}.",
                      implicacion="Ahí se concentra la necesidad documentada: es donde más se justifica invertir con medidas municipales.",
                      fuente="Ficha_dimensiones [OFICIAL]"))
    ca = ficha[(ficha.componente == "Capacidad adaptativa") & (ficha.origen != "NO DISPONIBLE")]
    for x in MUNICIPIOS:
        s = ca[ca.municipio == x]
        if len(s) >= 2 and s.valor.max() - s.valor.min() >= 0.2:
            hi, lo = s.loc[s.valor.idxmax()], s.loc[s.valor.idxmin()]
            H.append(dict(hallazgo=f"Tener capacidad en una dimensión no significa tenerla en todas: {x} tiene CA {fmt(hi.valor)} en "
                                   f"{hi.dimension.lower()} pero {fmt(lo.valor)} en {lo.dimension.lower()}.",
                          implicacion="Las medidas que fortalecen capacidad rinden más donde la brecha es mayor.",
                          fuente="Ficha_dimensiones [OFICIAL]"))
            break
    ev = d["evidencia"].groupby("ID").Evidencia.mean()
    top_ids = ev[ev == ev.max()].index.tolist()
    partes = []
    for mid in top_ids:
        nombre = d["medidas"][d["medidas"].ID == mid].Medida.iloc[0]
        f_ = res["frecuencia_por_medida"].get(mid, 0)
        en = "está en el portafolio" if mid in set(port.id) else "no entra en el portafolio"
        partes.append(f"'{nombre}' sale en el {f_:.0%} de los escenarios y {en}")
    H.append(dict(hallazgo=("Las medidas con más respaldo en los estudios CORNARE (evidencia en los tres municipios) son: "
                            + "; ".join(d['medidas'][d['medidas'].ID == m].Medida.iloc[0] for m in top_ids) + "."),
                  implicacion="Respaldo documental ≠ selección automática: " + "; ".join(partes) + ".",
                  fuente="Evidencia (Informes 2012/2014) [DERIVADO]"))
    fut = []
    for x in MUNICIPIOS:
        for dim in DIMENSIONES:
            r0, _ = valor(ficha, x, dim, "Riesgo hoy"); r1, _ = valor(ficha, x, dim, "Riesgo SSP3-7.0 2060")
            if r0 is not None and r1 is not None and r1 > r0:
                fut.append(f"{dim.lower()} en {x} ({fmt(r0)} → {fmt(r1)})")
    n60 = (ficha[ficha.componente == "Riesgo SSP3-7.0 2060"].origen != "NO DISPONIBLE").sum()
    H.append(dict(hallazgo=(f"Hacia 2060 (SSP3-7.0) el riesgo sube en {', '.join(fut)}." if fut else "No hay datos de cambio hacia 2060.")
                           + f" Solo {n60} de 21 celdas 2060 tienen dato.",
                  implicacion=f"Stress test: el portafolio {res['veredicto_2060']}. La decisión robusta frente a hoy y 2060 se mantiene.",
                  fuente="Ficha_dimensiones [OFICIAL]; motor"))
    H.append(dict(hallazgo=f"Solo el {base_pct:.0%} de los valores de vulnerabilidad y capacidad adaptativa están disponibles.",
                  implicacion=f"La decisión se elige por robustez (3.000 escenarios) y se identifican {len(d['voi'])} datos que podrían cambiarla.",
                  fuente="1_preparar.py [CALCULADO]"))
    return pd.DataFrame(H[:5])


# ----------------------------------------------------------------------------------------------
# EXCEL
# ----------------------------------------------------------------------------------------------
def guardar_excel(ruta, hojas, total_presupuesto):
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre[:31], index=False)
    wb = load_workbook(ruta)
    H = Font(name="Arial", bold=True, color="FFFFFF", size=10); FILL = PatternFill("solid", fgColor="1F4E5A")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = H; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max(10, min(60, largo * 0.9))
            for c in col[1:]:
                c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
    ws = wb["Portafolio_final"]
    col_costo = [c.column_letter for c in ws[1] if c.value == "costo_M"][0]
    n = ws.max_row
    ws[f"A{n + 2}"] = "TOTAL asignado (COP millones)"; ws[f"A{n + 2}"].font = Font(name="Arial", bold=True, size=10)
    ws[f"{col_costo}{n + 2}"] = f"=SUM({col_costo}2:{col_costo}{n})"
    ws[f"A{n + 3}"] = "Presupuesto (COP millones)"; ws[f"{col_costo}{n + 3}"] = total_presupuesto
    ws[f"A{n + 4}"] = "Chequeo"; ws[f"{col_costo}{n + 4}"] = f'=IF({col_costo}{n + 2}<={col_costo}{n + 3},"OK: dentro del presupuesto","EXCEDE")'
    for r in (n + 2, n + 3, n + 4):
        ws[f"{col_costo}{r}"].font = Font(name="Arial", bold=True, size=10)
    wb.save(ruta)


# ----------------------------------------------------------------------------------------------
def main():
    d = cargar()
    os.makedirs(SAL, exist_ok=True)
    elegidas, decision_txt = portafolio_final(d)
    port = construir_portafolio(d, elegidas)
    residual = construir_residual(d, port)
    mea, base_pct = construir_mea(d, port)
    brechas = construir_brechas(d, base_pct)
    hallazgos = construir_hallazgos(d, port, residual, base_pct)
    P = d["par"]["presupuesto"]; total = float(port.costo_M.sum())
    por_muni = port.groupby("territorio").costo_M.sum().reindex(MUNICIPIOS + ["Corredor"]).fillna(0)

    port.to_csv(os.path.join(SAL, "portafolio_final.csv"), index=False, encoding="utf-8-sig")
    residual.to_csv(os.path.join(SAL, "riesgo_residual.csv"), index=False, encoding="utf-8-sig")
    mea.to_csv(os.path.join(SAL, "seguimiento_MEA.csv"), index=False, encoding="utf-8-sig")
    brechas.to_csv(os.path.join(SAL, "brechas_informacion.csv"), index=False, encoding="utf-8-sig")
    hallazgos.to_csv(os.path.join(SAL, "hallazgos.csv"), index=False, encoding="utf-8-sig")

    # ---- JSON para el tablero (Vue). Cada bloque trae su origen para la trazabilidad (módulo 13)
    res = d["resumen"]
    js = dict(
        meta=dict(titulo="Motor de decisión de adaptación — Rionegro · Guarne · Marinilla", presupuesto_M=P,
                  decision_stress=decision_txt, n_simulaciones=int(d["par"]["n_simulaciones"]),
                  etiquetas=["OFICIAL", "CALCULADO", "DERIVADO", "USUARIO", "SUPUESTO", "INFERENCIA", "PROPUESTO", "NO DISPONIBLE"]),
        kpis=dict(presupuesto_M=P, asignado_M=total, saldo_M=P - total, n_intervenciones=len(port),
                  por_territorio_M=por_muni.to_dict(), cobertura_datos_VCA=round(base_pct, 3),
                  veredicto_2060=res["veredicto_2060"]),
        hallazgos=hallazgos.to_dict(orient="records"),
        portafolio=port.to_dict(orient="records"),
        escenarios=dict(referencia=res["referencia"], stress_2060=res["stress_2060"], robusto=res["robusto"],
                        entra_2060=res["entra_2060"], sale_2060=res["sale_2060"], aguante_robustez=res["aguante_robustez"],
                        portafolios_equivalentes=res["portafolios_equivalentes"]),
        frecuencia_por_medida=res["frecuencia_por_medida"],
        alternativas=d["alts"].where(pd.notna(d["alts"]), None).to_dict(orient="records"),
        ficha=d["ficha"].where(pd.notna(d["ficha"]), None).to_dict(orient="records"),
        celdas_vacias=d["vacias"].to_dict(orient="records"),
        datos_que_cambian_la_decision=d["voi"].where(pd.notna(d["voi"]), None).to_dict(orient="records") if len(d["voi"]) else [],
        contrafactuales=d["cf"].where(pd.notna(d["cf"]), None).to_dict(orient="records"),
        regret_top=d["regret"].head(10).to_dict(orient="records"),
        riesgo_residual=residual.to_dict(orient="records"),
        seguimiento_mea=mea.to_dict(orient="records"),
        brechas=brechas.to_dict(orient="records"),
        pesos=d["pesos"].to_dict(orient="records"),
        medidas=d["medidas"][["ID", "Dimensión", "Medida", "Costo", "Escala", "Actúa sobre"]].rename(
            columns={"Dimensión": "dimension", "Actúa sobre": "actua_sobre"}).to_dict(orient="records"),
    )
    def limpio(o):                                   # JSON válido para el navegador: NaN → null, numpy → Python
        if isinstance(o, dict):
            return {str(k): limpio(v) for k, v in o.items()}
        if isinstance(o, (list, tuple)):
            return [limpio(v) for v in o]
        if hasattr(o, "item"):
            o = o.item()
        if isinstance(o, float) and (o != o or o in (float("inf"), float("-inf"))):
            return None
        return o
    json.dump(limpio(js), open(os.path.join(SAL, "resultados.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1,
              allow_nan=False)
    # copia para el tablero Vue (si la carpeta frontend existe junto a los scripts)
    pub = os.path.join(AQUI, "frontend", "public")
    if os.path.isdir(os.path.join(AQUI, "frontend")):
        os.makedirs(pub, exist_ok=True)
        import shutil
        shutil.copy(os.path.join(SAL, "resultados.json"), os.path.join(pub, "resultados.json"))

    guardar_excel(os.path.join(SAL, "reporte_completo.xlsx"), {
        "Hallazgos": hallazgos, "Portafolio_final": port, "Riesgo_residual": residual, "Seguimiento_MEA": mea,
        "Brechas_informacion": brechas, "Alternativas": d["alts"], "Datos_que_cambian": d["voi"],
        "Contrafactuales": d["cf"], "Regret": d["regret"].head(30), "Ficha": d["ficha"], "Celdas_vacias": d["vacias"]}, P)

    # ---- pantalla
    L = "=" * 86
    print(L + "\n PASO 5 · REPORTES\n" + L)
    print(f"\nDecisión del stress test: {decision_txt}")
    print("\nHALLAZGOS (Producto 1)")
    for i, h in enumerate(hallazgos.itertuples(), 1):
        print(f"  {i}. {h.hallazgo}\n     → {h.implicacion}")
    print("\nPORTAFOLIO FINAL (Producto 2)")
    for r in port.itertuples():
        print(f"\n  {r.orden}. {r.medida}  ·  {r.territorio}  ·  ${r.costo_M:,.0f} M  ·  confianza {r.confianza}")
        print(f"     Problema:   {r.problema}")
        print(f"     Reduce:     {r.sensibilidad_que_reduce}  |  Fortalece: {r.capacidad_que_fortalece}")
        print(f"     Dónde:      {r.localizacion[:160]}{'…' if len(r.localizacion) > 160 else ''}")
        print(f"     Actores:    {r.actores}")
        print(f"     Por qué:    {r.explicacion}")
        print(f"     2060:       {r.cambio_ssp370_2060}")
    print(f"\n  TOTAL ${total:,.0f} M de ${P:,.0f} M ({'OK' if total <= P else 'EXCEDE'}) · "
          + " · ".join(f"{k} ${v:,.0f} M" for k, v in por_muni.items()))
    print("\nRIESGO RESIDUAL (Producto 3) — prioridad ALTA  (MEDIA y DESCONOCIDA en riesgo_residual.csv)")
    print("  Conteo:", residual.prioridad_residual.value_counts().to_dict())
    for r in residual[residual.prioridad_residual.isin(["ALTA"])].itertuples():
        print(f"  [ALTA] {r.municipio:9s} · {r.dimension:20s} · {r.estado_intervencion}")
        if r.problemas_documentados_sin_medida != "—":
            print(f"         sin medida: {r.problemas_documentados_sin_medida[:150]}{'…' if len(r.problemas_documentados_sin_medida) > 150 else ''}")
    desc = residual[residual.prioridad_residual == "DESCONOCIDA"]
    print(f"  [DESCONOCIDA] {len(desc)} combinaciones municipio × dimensión sin datos ni medida "
          f"(no se puede afirmar que su riesgo sea bajo)")
    print("\nSEGUIMIENTO MEA (Producto 3)")
    for r in mea.itertuples():
        print(f"  {r.orden}. {r.alternativa}: actividad → {r.actividad_que_se_hizo[:70]} | resultado → {r.resultado_que_cambio[:70]}")
    print(f"\nBRECHAS DE INFORMACIÓN: {len(brechas)} filas (ver brechas_informacion.csv)")
    print(f"\nOK · archivos en {os.path.relpath(SAL, AQUI)}/ : portafolio_final.csv, riesgo_residual.csv, seguimiento_MEA.csv, "
          f"brechas_informacion.csv, hallazgos.csv, resultados.json, reporte_completo.xlsx")


if __name__ == "__main__":
    main()
