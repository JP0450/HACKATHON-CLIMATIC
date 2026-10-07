# -*- coding: utf-8 -*-
"""
PASO 2 · PREPARAR DATOS  (módulos 1–4: ingesta, calidad, matriz de riesgo, catálogo de medidas)

Lee  datos/datos_entrada.xlsx  (la plantilla que edita el usuario) y deja tablas limpias en  salidas/preparado/
para que el motor (2_motor.py) las use. No inventa ni corrige datos: solo valida, clasifica y reporta.

Uso:   python 1_preparar.py
"""
import json
import os
import sys

import pandas as pd

try:
    sys.stdout.reconfigure(encoding="utf-8")   # para que Windows muestre bien los símbolos
except Exception:
    pass
from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
ENTRADA = os.path.join(AQUI, "datos", "datos_entrada.xlsx")
HISTORICO = os.path.join(AQUI, "datos", "fuentes", "REPORTE_MEDIDAS_ADAPTACION_MUNICIPIOS.xlsx")
SALIDA = os.path.join(AQUI, "salidas", "preparado")

MUNICIPIOS = ["Rionegro", "Guarne", "Marinilla"]
DIMENSIONES = ["Biodiversidad", "Recurso hídrico", "Seguridad alimentaria", "Hábitat",
               "Infraestructura", "Riesgo de desastres", "Salud"]
COMPONENTES = ["Amenaza", "Sensibilidad", "Capacidad adaptativa", "Vulnerabilidad",
               "Riesgo hoy", "Riesgo SSP3-7.0 2060"]
CRITERIOS = ["Necesidad", "Brecha CA", "Evidencia", "Cobertura", "Cobeneficios", "Robustez 2060"]

errores, avisos = [], []


# ----------------------------------------------------------------------------------------------
# MÓDULO 1 · INGESTA: abrir la plantilla
# ----------------------------------------------------------------------------------------------
def abrir_plantilla():
    if not os.path.exists(ENTRADA):
        sys.exit(f"No se encontró {ENTRADA}")
    wb = load_workbook(ENTRADA)                      # valores tal como los escribió el usuario
    hojas = ["Ficha_dimensiones", "Promedios_regionales", "Medidas", "Evidencia", "Pesos", "Parametros"]
    faltan = [h for h in hojas if h not in wb.sheetnames]
    if faltan:
        sys.exit(f"Faltan hojas en la plantilla: {faltan}")
    return wb


def origen_de(celda):
    """Clasifica una celda: OFICIAL / CALCULADO / USUARIO / NO DISPONIBLE (según su comentario)."""
    if celda.value is None or str(celda.value).strip() == "":
        return "NO DISPONIBLE"
    nota = celda.comment.text if celda.comment else ""
    if "[OFICIAL]" in nota:
        return "OFICIAL"
    if "[CALCULADO]" in nota:
        return "CALCULADO"
    return "USUARIO"


def como_numero(v, entero_grande=False):
    """Convierte a número. Con entero_grande=True acepta formatos como '6.000', '6,000', '$6.000 M' (presupuesto)."""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    txt = str(v).strip().replace("$", "").replace("M", "").replace(" ", "")
    if entero_grande:
        txt = txt.replace(".", "").replace(",", "")      # separadores de miles: 6.000 → 6000
    else:
        txt = txt.replace(",", ".")                      # decimal con coma: 0,61 → 0.61
    try:
        return float(txt)
    except ValueError:
        return "NO_NUMERICO"


# ----------------------------------------------------------------------------------------------
# MÓDULO 2 y 3 · CALIDAD + MATRIZ DE RIESGO: la ficha municipio × dimensión
# ----------------------------------------------------------------------------------------------
def leer_ficha(wb):
    ws = wb["Ficha_dimensiones"]
    encabezado = [c.value for c in ws[1]]
    col = {nombre: i for i, nombre in enumerate(encabezado)}
    filas = []
    vistos = set()
    for fila in ws.iter_rows(min_row=2):
        muni, dim = fila[col["Municipio"]].value, fila[col["Dimensión"]].value
        if muni is None and dim is None:
            continue
        if muni not in MUNICIPIOS or dim not in DIMENSIONES:
            continue                                   # notas al pie de la hoja
        if (muni, dim) in vistos:
            errores.append(f"Ficha: fila duplicada {muni} · {dim}")
        vistos.add((muni, dim))
        fuente_usuario = fila[col["Fuente de datos ingresados por el usuario"]].value
        for comp in COMPONENTES:
            celda = fila[col[comp]]
            valor = como_numero(celda.value)
            origen = origen_de(celda)
            if valor == "NO_NUMERICO":
                errores.append(f"Ficha: {muni} · {dim} · {comp} no es un número ('{celda.value}')")
                valor, origen = None, "NO DISPONIBLE"
            elif valor is not None and not (0 <= valor <= 1):
                errores.append(f"Ficha: {muni} · {dim} · {comp} = {valor} fuera del rango 0–1")
            if origen == "USUARIO" and not fuente_usuario:
                avisos.append(f"Ficha: {muni} · {dim} · {comp} fue ingresado sin anotar la fuente")
            fuente = (celda.comment.text.split("] ", 1)[-1] if origen in ("OFICIAL", "CALCULADO") and celda.comment
                      else (fuente_usuario or "") if origen == "USUARIO" else "")
            filas.append(dict(municipio=muni, dimension=dim, componente=comp, valor=valor, origen=origen,
                              fuente=fuente.replace("\n", " ")))
    for m in MUNICIPIOS:
        for d in DIMENSIONES:
            if (m, d) not in vistos:
                errores.append(f"Ficha: falta la fila {m} · {d}")
    return pd.DataFrame(filas)


def leer_promedios(wb):
    ws = wb["Promedios_regionales"]
    filas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila[0] in DIMENSIONES and fila[1] is not None:
            filas.append(dict(dimension=fila[0], v_promedio=float(fila[1]), fuente=fila[2]))
    return pd.DataFrame(filas)


# ----------------------------------------------------------------------------------------------
# MÓDULO 4 · CATÁLOGO DE MEDIDAS + evidencia
# ----------------------------------------------------------------------------------------------
def leer_medidas(wb):
    df = pd.read_excel(ENTRADA, sheet_name="Medidas")
    df = df[df["ID"].astype(str).str.match(r"^M\d\d$")].copy()
    df = df.rename(columns={"Costo de referencia (COP millones)": "Costo"})
    if df["ID"].duplicated().any():
        errores.append(f"Medidas: IDs duplicados {df.loc[df['ID'].duplicated(), 'ID'].tolist()}")
    for _, m in df.iterrows():
        if pd.isna(m["Costo"]) or m["Costo"] <= 0:
            errores.append(f"Medidas: {m['ID']} sin costo válido")
        if m["Escala"] not in ("municipal", "corredor"):
            errores.append(f"Medidas: {m['ID']} escala '{m['Escala']}' (debe ser municipal o corredor)")
        if m["Actúa sobre"] not in ("S", "CA", "S+CA"):
            errores.append(f"Medidas: {m['ID']} 'Actúa sobre' = '{m['Actúa sobre']}' (S, CA o S+CA)")
        if m["Dimensión"] not in DIMENSIONES:
            errores.append(f"Medidas: {m['ID']} dimensión desconocida '{m['Dimensión']}'")
        for d in str(m["Cobeneficios (otras dimensiones)"]).split(";"):
            d = d.strip()
            if d and d != "nan" and d not in DIMENSIONES:
                errores.append(f"Medidas: {m['ID']} cobeneficio desconocido '{d}'")
    df["Cobeneficios (otras dimensiones)"] = df["Cobeneficios (otras dimensiones)"].fillna("")
    return df


def leer_evidencia(wb, medidas):
    df = pd.read_excel(ENTRADA, sheet_name="Evidencia")
    df = df[df["ID"].astype(str).str.match(r"^M\d\d$")].copy()
    df = df.rename(columns={"Evidencia (0 / 0,5 / 1)": "Evidencia"})
    bad = df[~df["Evidencia"].isin([0, 0.5, 1])]
    for _, r in bad.iterrows():
        errores.append(f"Evidencia: {r['ID']} · {r['Municipio']} = {r['Evidencia']} (debe ser 0, 0,5 o 1)")
    for mid in medidas["ID"]:
        for muni in MUNICIPIOS:
            n = ((df["ID"] == mid) & (df["Municipio"] == muni)).sum()
            if n == 0:
                errores.append(f"Evidencia: falta {mid} · {muni}")
            elif n > 1:
                errores.append(f"Evidencia: {mid} · {muni} está repetida")
    return df


def leer_pesos(wb):
    ws = wb["Pesos"]
    base, rob = {}, None
    for fila in ws.iter_rows(min_row=2, max_row=12):
        nombre = fila[0].value
        if nombre in CRITERIOS:
            base[nombre] = como_numero(fila[3].value)
        if nombre == "Peso de robustez en el stress test":
            rob = como_numero(fila[1].value)
    if set(base) != set(CRITERIOS) or any(v is None or v == "NO_NUMERICO" for v in base.values()):
        errores.append(f"Pesos: faltan pesos o no son números: {base}")
        return None
    if rob is None or rob == "NO_NUMERICO" or not (0 <= rob < 1):
        errores.append("Pesos: el peso de robustez del stress test debe estar entre 0 y 1")
        return None
    if abs(sum(base.values()) - 1) > 1e-6:
        errores.append(f"Pesos: los pesos de referencia suman {sum(base.values()):.1%} (deben sumar 100%)")
    # Stress test: robustez = rob; los demás bajan en la misma proporción (se recalcula aquí, no se lee de la fórmula)
    otros = {k: v for k, v in base.items() if k != "Robustez 2060"}
    s = sum(otros.values())
    stress = {k: v / s * (1 - rob) for k, v in otros.items()}
    stress["Robustez 2060"] = rob
    return pd.DataFrame({"criterio": CRITERIOS,
                         "referencia": [base[c] for c in CRITERIOS],
                         "stress_ssp370_2060": [stress[c] for c in CRITERIOS]})


def leer_parametros(wb):
    ws = wb["Parametros"]
    claves = {"Presupuesto (COP millones)": "presupuesto",
              "Umbral de evidencia mínima": "umbral_evidencia",
              "Alineación de medidas que solo reducen sensibilidad": "alineacion_S",
              "Valor de referencia para celdas vacías": "valor_vacio_base",
              "Elasticidad beneficio–costo (α)": "alfa",
              "Nº de simulaciones": "n_simulaciones",
              "Semilla aleatoria": "semilla"}
    p = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila[0] == "Decisión del stress test (A / B / C)":
            d = str(fila[1] or "A").strip().upper()
            if d not in ("A", "B", "C"):
                errores.append(f"Parámetros: decisión del stress test '{fila[1]}' (debe ser A, B o C)")
            p["decision_stress"] = d
            continue
        if fila[0] in claves:
            v = como_numero(fila[1], entero_grande=(claves[fila[0]] in ("presupuesto", "n_simulaciones", "semilla")))
            if v is None or v == "NO_NUMERICO":
                errores.append(f"Parámetros: '{fila[0]}' vacío o no numérico")
            p[claves[fila[0]]] = v
    p.setdefault("decision_stress", "A")
    for k in claves.values():
        if k not in p:
            errores.append(f"Parámetros: falta '{k}'")
    if isinstance(p.get("presupuesto"), float) and p["presupuesto"] > 100000:
        errores.append(f"Parámetros: presupuesto = {p['presupuesto']:,.0f}. Escríbalo en MILLONES de pesos "
                       f"(5.000 millones → 5000), no en pesos")
    if isinstance(p.get("presupuesto"), float) and p["presupuesto"] < 600:
        errores.append(f"Parámetros: presupuesto = {p['presupuesto']:,.0f} M. Debe estar en millones de pesos "
                       f"(ej. 5000) y ser al menos 600 (la medida más barata)")
    for k in ("umbral_evidencia", "alineacion_S", "valor_vacio_base", "alfa"):
        if isinstance(p.get(k), float) and not (0 <= p[k] <= 1):
            errores.append(f"Parámetros: '{k}' = {p[k]} debe estar entre 0 y 1")
    return p


# ----------------------------------------------------------------------------------------------
# Antecedentes: registro histórico de medidas de adaptación (solo contexto, no entra al puntaje)
# ----------------------------------------------------------------------------------------------
def resumir_historico():
    if not os.path.exists(HISTORICO):
        avisos.append("No se encontró el Excel histórico de medidas; se omite el resumen de antecedentes")
        return None, None
    h = pd.read_excel(HISTORICO)
    calidad = dict(filas=len(h), filas_duplicadas=int(h.duplicated().sum()),
                   inversion_vacia=int(h["Valor Inversion"].isna().sum()),
                   ecosistema_prueba=int((h["Ecosistema"].astype(str).str.lower().str.contains("prueba")).sum()),
                   vulnerabilidad_no_aplica=int((h["Vulnerabilidad"] == "No Aplica").sum()))
    corr = h[h["Municipio"].str.upper().isin([m.upper() for m in MUNICIPIOS])].copy()
    corr["Municipio"] = corr["Municipio"].str.title()
    return corr, calidad


# ----------------------------------------------------------------------------------------------
# Reporte de calidad
# ----------------------------------------------------------------------------------------------
def calidad_por_componente(ficha):
    t = ficha.pivot_table(index="componente", columns="origen", values="valor", aggfunc="size", fill_value=0)
    for c in ["OFICIAL", "CALCULADO", "USUARIO", "NO DISPONIBLE"]:
        if c not in t.columns:
            t[c] = 0
    t = t[["OFICIAL", "CALCULADO", "USUARIO", "NO DISPONIBLE"]].reindex(COMPONENTES)
    t["con_dato_%"] = ((t["OFICIAL"] + t["CALCULADO"] + t["USUARIO"]) / 21 * 100).round(0)
    # [SUPUESTO] escala de calidad: Alta ≥ 75% de celdas con dato, Media ≥ 25%, Baja < 25%
    t["calidad"] = pd.cut(t["con_dato_%"], bins=[-1, 24.99, 74.99, 101], labels=["Baja", "Media", "Alta"])
    return t


def main():
    print("=" * 78)
    print(" PASO 2 · PREPARAR DATOS")
    print("=" * 78)
    wb = abrir_plantilla()
    ficha = leer_ficha(wb)
    promedios = leer_promedios(wb)
    medidas = leer_medidas(wb)
    evidencia = leer_evidencia(wb, medidas)
    pesos = leer_pesos(wb)
    parametros = leer_parametros(wb)
    historico, cal_hist = resumir_historico()

    os.makedirs(SALIDA, exist_ok=True)
    ficha.to_csv(os.path.join(SALIDA, "ficha_larga.csv"), index=False, encoding="utf-8-sig")
    promedios.to_csv(os.path.join(SALIDA, "promedios_regionales.csv"), index=False, encoding="utf-8-sig")
    medidas.to_csv(os.path.join(SALIDA, "medidas.csv"), index=False, encoding="utf-8-sig")
    evidencia.to_csv(os.path.join(SALIDA, "evidencia.csv"), index=False, encoding="utf-8-sig")
    if pesos is not None:
        pesos.to_csv(os.path.join(SALIDA, "pesos.csv"), index=False, encoding="utf-8-sig")
    with open(os.path.join(SALIDA, "parametros.json"), "w", encoding="utf-8") as fh:
        json.dump(parametros, fh, ensure_ascii=False, indent=1)
    if historico is not None:
        historico.to_csv(os.path.join(SALIDA, "historico_corredor.csv"), index=False, encoding="utf-8-sig")

    # ---- resumen en pantalla
    cal = calidad_por_componente(ficha)
    print("\n1) Ficha municipio × dimensión (21 filas × 6 componentes = 126 celdas)")
    print(cal.to_string())
    n = ficha["origen"].value_counts()
    print(f"\n   Totales: OFICIAL {n.get('OFICIAL', 0)} · CALCULADO {n.get('CALCULADO', 0)} · "
          f"USUARIO {n.get('USUARIO', 0)} · NO DISPONIBLE {n.get('NO DISPONIBLE', 0)}")
    vca = ficha[ficha.componente.isin(["Vulnerabilidad", "Capacidad adaptativa"])]
    con = (vca.origen != "NO DISPONIBLE").sum()
    print(f"   Vulnerabilidad + Capacidad adaptativa (lo que usa el puntaje): {con} de {len(vca)} celdas con dato "
          f"({con / len(vca):.0%})")

    print("\n2) Datos con valor (lo demás es NO DISPONIBLE y queda como entrada del usuario):")
    for _, r in ficha[ficha.origen != "NO DISPONIBLE"].iterrows():
        print(f"   [{r.origen:9s}] {r.municipio:9s} · {r.dimension:20s} · {r.componente:22s} = {r.valor:.3f}")

    print(f"\n3) Medidas: {len(medidas)} · costo total si se hicieran todas una vez: "
          f"${medidas['Costo'].sum():,.0f} M (presupuesto ${parametros.get('presupuesto', 0):,.0f} M)")
    print(f"   Evidencia: {len(evidencia)} calificaciones · con evidencia (≥0,5): {(evidencia.Evidencia >= 0.5).sum()} · "
          f"sin evidencia: {(evidencia.Evidencia == 0).sum()}")
    if pesos is not None:
        print("\n4) Pesos")
        print(pesos.assign(referencia=lambda d: (d.referencia * 100).round(2).astype(str) + "%",
                           stress_ssp370_2060=lambda d: (d.stress_ssp370_2060 * 100).round(3).astype(str) + "%").to_string(index=False))
    print(f"\n5) Parámetros: {parametros}")

    if historico is not None:
        print("\n6) Antecedentes (Excel histórico de medidas de adaptación, 2018–2025) — solo contexto")
        print(f"   Calidad del archivo: {cal_hist}")
        g = historico.groupby("Municipio")["Valor Inversion"].agg(registros="count", inversion_COP="sum")
        g = g.reindex(MUNICIPIOS).fillna(0)
        g["inversion_COP"] = g["inversion_COP"].map(lambda v: f"${v / 1e6:,.1f} M")
        g = g.rename(columns={"inversion_COP": "inversión histórica (COP millones)"})
        print(g.to_string())
        print("   Nota: pocos registros (Marinilla) indican posible subregistro, NO menor necesidad.")

    print("\n" + "-" * 78)
    for a in avisos:
        print("AVISO:", a)
    if errores:
        print(f"\n{len(errores)} ERROR(ES) — corrija datos/datos_entrada.xlsx y vuelva a ejecutar:")
        for e in errores:
            print("  ✗", e)
        sys.exit(1)
    print(f"OK · datos listos en {os.path.relpath(SALIDA, AQUI)}/  → siguiente paso: python 2_motor.py")


if __name__ == "__main__":
    main()
